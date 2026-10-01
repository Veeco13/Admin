# -*- coding: utf-8 -*-
"""
الخطابات والشهادات (القسم 26)

- شهادة راتب «إلى من يهمه الأمر» و«استمرارية راتب»: من قالب الشركة نفسه (forms/salary_certificate.xlsx — بيتطبع على
  الورق الرسمي، والشكل مابيتغيّرش): البيانات الوظيفية للموظف بس اللي بتتملى، وبعدين PDF للمعاينة والطباعة.
  البدلات مابتتذكرش في أي مكان (استمرارية الراتب: سطر «الراتب» بس).
  التواريخ بتتكتب نص ثابت باسم الشهر («11-Oct-2023» / «11 أكتوبر 2023») عشان مايبقاش فيه لبس بين اليوم والشهر.
- نموذج الإجازة: بيتطبع من المتصفح (تصميم جديد)، والطلب بيتحفظ بحالته (مقدَّم / معتمد / مرفوض) لمرحلة الإجازات بعدين.
- نموذج العودة من الإجازة: مربوط بطلب الإجازة (أو بتواريخ إجازة مش متسجّلة)، وبيحسب التأخير عن تاريخ العودة المقرر.
- كل خطاب ليه رقم (HR-SCR-2026-0001 للشهادات، LV-2026-0001 للإجازة، RT-2026-0001 للعودة) ونسخة من بياناته وقت إصداره (data)،
  فإعادة الطباعة بتطلع نفس الخطاب بنفس الرقم حتى لو بيانات الموظف اتغيّرت بعدها.
"""
import io
import json
import math
import os
import re
from copy import copy
from datetime import date, timedelta

import openpyxl
from sqlalchemy import func, select

import custody_excel
import db
import models as M
import value_i18n

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "forms", "salary_certificate.xlsx")
KINDS = {"salary": ("SCR", "شهادة راتب"), "continuity": ("SCR", "شهادة استمرارية راتب"), "leave": ("LV", "طلب إجازة"),
         "return": ("RT", "عودة من إجازة")}
SHEETS = {"salary": "شهادة راتب", "continuity": "استمرارية راتب"}
FORMS = ("leave", "return")          # نماذج بتتطبع من المتصفح (صلاحية «تعديل الموظفين» ومن غير «المرتب»)
PREFIX = {"SCR": "HR-SCR", "LV": "LV", "RT": "RT"}
MONTHS_EN = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
MONTHS_AR = ("يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر")
LEAVE_TYPES = {"paid": "مدفوعة", "unpaid": "بدون راتب", "advance": "مدفوعة مقدمًا", "with_salary": "مع الراتب",
               "rotation": "إجازة تناوب (Rotation)", "compassionate": "إجازة ظرف خاص (Compassionate)"}
LEAVE_STATUS = ("submitted", "approved", "rejected")

# البنوك في الكويت ← (عربي، إنجليزي، كلمات للتعرّف) — الأطول الأول عشان «الأهلي المتحد» قبل «الأهلي»
BANKS = [("البنك الأهلي المتحد", "Ahli United Bank", ("ahli united", "الاهلي المتحد", "المتحد")),
         ("بنك الكويت الوطني", "National Bank of Kuwait", ("national bank", "nbk", "الوطني")),
         ("بيت التمويل الكويتي", "Kuwait Finance House", ("finance house", "kfh", "بيت التمويل")),
         ("بنك برقان", "Burgan Bank", ("burgan", "برقان")),
         ("بنك الخليج", "Gulf Bank", ("gulf", "الخليج")),
         ("البنك التجاري الكويتي", "Commercial Bank of Kuwait", ("commercial", "cbk", "التجاري")),
         ("بنك بوبيان", "Boubyan Bank", ("boubyan", "بوبيان")),
         ("بنك وربة", "Warba Bank", ("warba", "وربة", "وربه")),
         ("بنك الكويت الدولي", "Kuwait International Bank", ("international", "kib", "الدولي")),
         ("البنك الأهلي الكويتي", "Al Ahli Bank of Kuwait", ("ahli", "abk", "الاهلي", "الأهلي"))]


def _norm(v):
    v = re.sub(r"[ً-ْ]", "", str(v or "")).strip().lower()
    return v.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ى", "ي")


def bank_names(text):
    """اسم البنك زي ما هو متسجّل ← (عربي، إنجليزي). لو مش معروف بيرجع نفس الاسم في الاتنين."""
    n = _norm(text)
    if not n:
        return "", ""
    for ar, en, keys in BANKS:
        if any(_norm(k) in n for k in keys):
            return ar, en
    return text, text


def words_en(amount):
    """800 ← «(Eight Hundred KWD Only)» زي القالب."""
    kd, fils = divmod(round(float(amount or 0) * 1000), 1000)
    text = custody_excel._words(kd).title().replace(" And ", " and ")
    return f"({text} KWD{f' and {fils} Fils' if fils else ''} Only)"


def words_ar(amount):
    """800 ← «( ثمانمائة دينار كويتي لا غير )» زي القالب."""
    text = re.sub(r"^فقط\s+", "", custody_excel.amount_words_ar(amount))
    return f"( {text} )"


def _primary_company(s, e):
    aff = s.scalar(select(M.EmployeeAffiliation).where(M.EmployeeAffiliation.employeeId == e.id)
                   .order_by(M.EmployeeAffiliation.position).limit(1))
    return s.get(M.Company, aff.companyId) if aff is not None and aff.companyId else None


def employee_basics(s, e):
    """بيانات الموظف المشتركة في كل الخطابات (بالعربي والإنجليزي)."""
    tr = value_i18n.merged(s)
    co = _primary_company(s, e)
    return {
        "employeeId": e.id, "nameAr": e.name or "", "nameEn": e.nameEn or "",
        "nationalityAr": e.nationality or "", "nationalityEn": e.nationalityEn or value_i18n.lookup(tr, "nationality", e.nationality) or "",
        "jobAr": e.profession or "", "jobEn": e.professionEn or value_i18n.lookup(tr, "profession", e.profession) or "",
        "civilId": e.id, "hireDate": db.ser(e.dateOfHire),
        "companyId": co.id if co is not None else None, "companyAr": co.nameAr if co is not None else "",
        "companyEn": (co.nameEn or co.nameAr) if co is not None else "",
    }


def salary_defaults(s, e, see_bank):
    """القيم الافتراضية لشهادة الراتب (بتتعرض في النافذة، والوظيفة والبنك والجهة بتتعدّل)."""
    d = employee_basics(s, e)
    bank_ar, bank_en = bank_names(e.bank) if see_bank else ("", "")
    salary = float(e.salary or 0)
    d.update({"salary": salary, "wordsAr": words_ar(salary) if salary else "", "wordsEn": words_en(salary) if salary else "",
              "bankAr": bank_ar, "bankEn": bank_en, "account": (e.iban or "") if see_bank else "", "toAr": "", "toEn": "",
              "date": date.today().isoformat()})
    return d


def next_number(s, kind, year):
    series = KINDS[kind][0]
    no = (s.scalar(select(func.max(M.HrLetter.no)).where(M.HrLetter.series == series, M.HrLetter.year == year)) or 0) + 1
    return series, no, f"{PREFIX[series]}-{year}-{no:04d}"


def _leave_dates(d):
    """نوع الإجازة وتواريخها من الطلب ← (النوع، البداية، النهاية) أو رسالة خطأ."""
    typ = d.get("leaveType")
    if typ not in LEAVE_TYPES:
        return None, "اختار نوع الإجازة"
    start, end = db.parse_date(d.get("from")), db.parse_date(d.get("to"))
    if not start or not end:
        return None, "تاريخ بداية ونهاية الإجازة مطلوبين"
    if end < start:
        return None, "نهاية الإجازة قبل بدايتها"
    return (typ, start, end), None


def _form_person(s, e):
    """بيانات الموظف في نماذج الإجازة والعودة (الإدارة، الرقم الوظيفي، الإقامة، عربياته)."""
    snap = employee_basics(s, e)
    cc = s.scalar(select(M.CostCenter).where(M.CostCenter.name == e.costCenter)) if e.costCenter else None
    snap.update({
        "department": e.costCenter or "", "departmentEn": (cc.nameEn if cc is not None else "") or e.costCenter or "",
        "employeeCode": e.dpId or "", "residencyExp": db.ser(e.residencyExp),
        "plates": [v.plate for v in s.scalars(select(M.Vehicle).where(M.Vehicle.driverId == e.id))],
        "date": date.today().isoformat(),
    })
    return snap


def leave_snapshot(s, e, d):
    """طلب إجازة ← بياناته (الموظف + تفاصيل الإجازة) أو رسالة خطأ."""
    got, msg = _leave_dates(d)
    if msg:
        return None, msg
    typ, start, end = got
    snap = _form_person(s, e)
    snap.update({
        "leaveType": typ, "from": start.isoformat(), "to": end.isoformat(), "days": (end - start).days + 1,
        "returnDate": (end + timedelta(days=1)).isoformat(), "phone": (d.get("phone") or e.phone or "").strip(),
        "address": (d.get("address") or "").strip(), "notes": (d.get("notes") or "").strip(),
    })
    return snap, None


def return_of(s, leave_id, employee_id):
    """رقم نموذج العودة المسجّل لطلب الإجازة ده (لو فيه)."""
    for x in s.scalars(select(M.HrLetter).where(M.HrLetter.kind == "return", M.HrLetter.employeeId == employee_id)):
        if json.loads(x.data or "{}").get("leaveId") == leave_id:
            return x.number
    return None


def return_snapshot(s, e, d):
    """عودة من إجازة ← بياناتها أو رسالة خطأ. مربوطة بطلب إجازة من السيستم، أو بتواريخ إجازة مش متسجّلة (ورق قديم)."""
    lv = s.get(M.HrLetter, d["leaveId"]) if d.get("leaveId") else None
    if d.get("leaveId"):
        if lv is None or lv.kind != "leave" or lv.employeeId != e.id:
            return None, "طلب الإجازة مش موجود"
        if lv.status == "rejected":
            return None, "طلب الإجازة ده مرفوض"
        taken = return_of(s, lv.id, e.id)
        if taken:
            return None, f"العودة من الإجازة دي متسجّلة قبل كده ({taken})"
        typ, start, end = json.loads(lv.data or "{}").get("leaveType"), lv.dateFrom, lv.dateTo
    else:
        got, msg = _leave_dates(d)
        if msg:
            return None, msg
        typ, start, end = got
    actual = db.parse_date(d.get("actualDate"))
    if not actual:
        return None, "تاريخ العودة الفعلي مطلوب"
    if actual <= start:
        return None, "تاريخ العودة لازم يكون بعد بداية الإجازة"
    if actual > date.today():
        return None, "تاريخ العودة الفعلي لسه ماجاش"
    due = end + timedelta(days=1)
    snap = _form_person(s, e)
    snap.update({
        "leaveId": lv.id if lv is not None else None, "leaveNumber": lv.number if lv is not None else "",
        "leaveType": typ, "from": start.isoformat(), "to": end.isoformat(), "days": (end - start).days + 1,
        "returnDate": due.isoformat(), "actualDate": actual.isoformat(), "delay": (actual - due).days,
        "reason": (d.get("reason") or "").strip(), "notes": (d.get("notes") or "").strip(),
    })
    return snap, None


def date_en(v):
    """2023-10-11 ← «11-Oct-2023» (نص ثابت — مابيعتمدش على إعدادات المنطقة في جهاز السيرفر)."""
    d = db.parse_date(v)
    return f"{d.day:02d}-{MONTHS_EN[d.month - 1]}-{d.year}" if d else None


def date_ar(v):
    """2023-10-11 ← «11 أكتوبر 2023»."""
    d = db.parse_date(v)
    return f"{d.day} {MONTHS_AR[d.month - 1]} {d.year}" if d else None


# ---------------------------------------------------------------------------
# ورق الشركة الرسمي (letterhead): صورة A4 لكل شركة. الطباعة بتبقى على الورق المطبوع نفسه، فالصورة بتظهر خلفية في
# **المعاينة بس**، ومقاساتها (آخر الترويسة وبداية التذييل) بتظبط هوامش الشهادة علشان مفيش حاجة تنزل على المطبوع.
# ---------------------------------------------------------------------------
PAGE_H_MM = 297.0
_PAPER_PDF = {}


def measure_paper(path):
    """صورة ورق الشركة ← (آخر الترويسة من فوق، بداية التذييل من تحت) بالملّي، أو None لو مش واضحة.
    أطول مساحة فاضية في الصفحة هي مكان الكتابة: اللي فوقها ترويسة واللي تحتها تذييل."""
    from PIL import Image
    try:
        with Image.open(path) as im:
            g = im.convert("L")
            w, h = g.size
            g = g.crop((int(w * .02), 0, int(w * .98), h))          # حواف المسح (ظل الورقة) مش حبر
            hist = g.histogram()
            paper = max(range(128, 256), key=lambda v: hist[v])     # لون الورق = أكتر درجة فاتحة
            ink = g.point(lambda v: 255 if v < paper - 20 else 0).resize((1, h), Image.BOX)   # نسبة الحبر في كل سطر
            rows = [v > 1 for v in ink.getdata()]
    except Exception:
        return None
    best, start = (0, 0), None
    for y, on in enumerate(rows + [True]):
        if not on and start is None:
            start = y
        elif on and start is not None:
            if y - start > best[1] - best[0]:
                best = (start, y)
            start = None
    if best[1] - best[0] < h * .4:                                  # مفيش مساحة كتابة واضحة ← مش ورق رسمي
        return None
    return round(best[0] / h * PAGE_H_MM, 1), round((h - best[1]) / h * PAGE_H_MM, 1)


def company_paper(s, company_id):
    """ورق الشركة ← {path, size: (فوق، تحت) أو None} — أو None لو مالهاش ورق مرفوع."""
    c = s.get(M.Company, company_id) if company_id else None
    path = db.resolve_file(c.letterheadPath) if c is not None and c.letterheadPath else None
    if not path or not os.path.exists(path):
        return None
    size = (c.letterheadTop, c.letterheadBottom) if c.letterheadTop is not None and c.letterheadBottom is not None else None
    return {"path": path, "size": size}


def paper_pdf(path):
    """صورة الورق ← PDF صفحة A4 واحدة (في الذاكرة لحد ما الملف يتغيّر)."""
    from PIL import Image
    key = (path, os.path.getmtime(path))
    if key not in _PAPER_PDF:
        with Image.open(path) as im:
            im = im.convert("RGB")
            im.thumbnail((1654, 2339))                              # 200 نقطة في البوصة كفاية للمعاينة
            buf = io.BytesIO()
            im.save(buf, "PDF", resolution=im.width / (210 / 25.4))
        if len(_PAPER_PDF) > 8:
            _PAPER_PDF.clear()
        _PAPER_PDF[key] = buf.getvalue()
    return _PAPER_PDF[key]


def with_paper(pdf, paper):
    """الشهادة فوق ورق الشركة — **للمعاينة بس** (الطباعة من الملف الأصلي على الورق الرسمي)."""
    from pypdf import PdfReader, PdfWriter, Transformation
    bg = PdfReader(io.BytesIO(paper)).pages[0]
    bw, bh = float(bg.mediabox.width), float(bg.mediabox.height)
    out = PdfWriter()
    for page in PdfReader(io.BytesIO(pdf)).pages:
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        base = out.add_blank_page(width=w, height=h)
        base.merge_transformed_page(bg, Transformation().scale(w / bw, h / bh))
        base.merge_page(page)
    buf = io.BytesIO()
    out.write(buf)
    return buf.getvalue()


def salary_xlsx(kind, number, data, paper=None):
    """قالب الشركة ← الشيت المطلوب بس، متعبّي بالبيانات الوظيفية (الشكل زي ما هو) ← xlsx bytes.
    paper = (آخر الترويسة، بداية التذييل) بالملّي من ورق الشركة ← الهوامش بتتظبط عليه."""
    wb = openpyxl.load_workbook(TEMPLATE)
    keep = SHEETS[kind]
    for ws in list(wb.worksheets):
        if ws.title != keep:
            wb.remove(ws)
    ws = wb[keep]

    def put(coord, v):
        ws[coord] = v
        # النص الطويل (زي اسم الشركة الإنجليزي) بيتقصّ أو يعدّي على الفاصل الرمادي ← الخط بيصغر على قد الخانة بس
        if isinstance(v, str) and coord[0] in "BCG":
            cell = ws[coord]
            size = cell.font.sz or 11
            cap = 36 * 11 / size * (1.4 if re.search(r"[؀-ۿ]", v) else 1)
            if len(v) > cap:
                f = copy(cell.font)
                f.sz = max(7, math.floor(size * cap / len(v) * 2) / 2)
                cell.font = f

    put("C9", date_en(data.get("date")))
    put("C11", data.get("toEn") or None)
    put("G11", data.get("toAr") or None)
    put("B13", data.get("companyEn"))
    put("G13", data.get("companyAr"))
    put("C16", data.get("nameEn"))
    put("G16", data.get("nameAr"))
    put("C17", data.get("nationalityEn"))
    put("G17", data.get("nationalityAr"))
    put("C18", data.get("civilId"))
    put("G18", data.get("civilId"))
    put("C19", date_en(data.get("hireDate")))
    put("G19", date_ar(data.get("hireDate")))
    al = copy(ws["G19"].alignment)
    al.readingOrder = 2          # من اليمين للشمال — غير كده «11 أكتوبر 2023» بيتقلب لـ «أكتوبر 2023 11»
    ws["G19"].alignment = al
    put("C20", data.get("jobEn"))
    put("G20", data.get("jobAr"))
    if kind == "salary":
        put("C21", data.get("salary"))
        put("G21", data.get("salary"))
        put("A22", data.get("wordsEn"))
        put("G22", data.get("wordsAr"))
        put("C23", data.get("bankEn") or None)
        put("G23", data.get("bankAr") or None)
        put("C24", data.get("account") or None)
        put("G24", data.get("account") or None)
    else:
        # استمرارية الراتب: سطر «الراتب» بس — البدلات والإجمالي مابيتذكروش
        put("A23", "Salary")
        put("I23", "الراتب")
        put("C23", data.get("salary"))
        put("G23", data.get("salary"))
        for r in (24, 25):
            for col in "ABCDEFGHI":
                ws[f"{col}{r}"].value = None
            ws.row_dimensions[r].hidden = True
        put("C28", data.get("bankEn") or None)
        put("G28", data.get("bankAr") or None)
        put("C29", data.get("account") or None)
        put("G29", data.get("account") or None)
    # رقم الشهادة: في تذييل الصفحة — أقصى الأسفل والشمال، بخط صغير (مش في نص الورقة)
    ws.oddFooter.left.text, ws.oddFooter.left.size, ws.oddFooter.left.font = f"Ref # {number}", 8, "Calibri,Regular"
    # منطقة الطباعة لحد آخر سطر مكتوب (السطور الفاضية تحت بتزق لصفحة تانية لما هامش التذييل يكبر)
    filled = [c.row for row in ws.iter_rows(min_row=2, max_row=45, max_col=9) for c in row if c.value not in (None, "", " ")]
    ws.print_area = f"A2:I{max(filled) + 1}"
    if paper:
        top, bottom = paper
        mg, default = ws.page_margins, ws.sheet_format.defaultRowHeight or 15
        blank = sum(ws.row_dimensions[r].height or default for r in range(2, min(filled))) * 25.4 / 72   # سطور فاضية قبل العنوان
        mg.top = max(mg.top, (top + 5 - blank) / 25.4)              # العنوان تحت الترويسة المطبوعة
        mg.footer = (bottom + 3) / 25.4                            # الرقم المرجعي فوق التذييل المطبوع
        mg.bottom = max(mg.bottom, mg.footer + 5 / 25.4)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_api(x, see_salary=True):
    d = {"id": x.id, "kind": x.kind, "number": x.number, "employeeId": x.employeeId, "companyId": x.companyId,
         "status": x.status, "dateFrom": db.ser(x.dateFrom), "dateTo": db.ser(x.dateTo), "days": x.days,
         "createdBy": x.createdBy, "createdAt": db.ser(x.createdAt), "decidedBy": x.decidedBy, "decidedAt": db.ser(x.decidedAt)}
    try:
        data = json.loads(x.data or "{}")
    except ValueError:
        data = {}
    if not see_salary:
        for k in ("salary", "wordsAr", "wordsEn", "bankAr", "bankEn", "account"):
            data.pop(k, None)
    d["data"] = data
    return d
