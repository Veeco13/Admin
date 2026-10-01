# -*- coding: utf-8 -*-
"""
كشوف العهد (openpyxl) بهوية التقارير — بتتعرض PDF للمعاينة والطباعة بس (contracts.xlsx_to_pdf)، مابتتنزّلش.
كلها باسم وشعار الشركة المُصدِرة (أبراج انرجي — custody.settings) لأن الفلوس منها وهي اللي بتفوتر مراكز التكلفة.

اللغة (lang): «ar» عربي بس (الشيت من اليمين للشمال، والمبلغ كتابةً بالعربي)، «en» إنجليزي بس، «both» ثنائي.
كل شيت شايل لغته (ws._lunx_lang)، والنصوص بتتكتب بـ _L(ws, إنجليزي، عربي) والأسماء بـ _V.

- تقفيل العهدة: شيت «الملخص» (الفواتير اللي طلعت + تسوية العهدة مع المستلم)، وشيت فاتورة لكل مركز تكلفة
  (INVOICE: موظف في كل صف وبند في كل عمود + الدعم الإداري، وملخص الفاتورة وتوقيعات الاعتماد)، وملحق اختياري
  بكشف فردي لكل موظف (كل كشف في صفحة مطبوعة لوحده).
- فاتورة لوحدها (invoice_workbook).
- طلب صرف عهدة (Advance Payment Request): كل الأشخاص وبنودهم + ملخص المطلوب لكل مركز تكلفة.

الأرقام معادلات (SUM). الطباعة: A4، عرض الصفحة، رأس الجدول بيتكرر، «صفحة X من Y»، وختم «صادر بواسطة المستخدم — الوقت» (من غير اسم البرنامج).
"""
import io
import math
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

import custody
import db
import models as M
import value_i18n

try:                                          # الشعار محتاج Pillow — من غيره الكشف بيطلع من غير شعار
    from openpyxl.drawing.image import Image as XLImage
    import PIL  # noqa: F401
except ImportError:                           # pragma: no cover
    XLImage = None

PRIMARY, PRIMARY_SOFT, ACCENT, INK, MUTED, LINE, ZEBRA, RED = (
    "1F5C4A", "E3EFEA", "C8963E", "17231F", "5F6F69", "D5DDDA", "F6F8F7", "C0392B")
FONT = "Arial"
KWD = '#,##0.000;-#,##0.000;"–"'
FORM_NO = ("Form No. T-104", "نموذج رقم T-104")
NOTE = ("Above charges are based on the actual announced charges of PAM, the Ministry of Interior, the Ministry of Health, "
        "the Ministry of Foreign Affairs and other authorities, and are subject to change if these authorities introduce or "
        "remove charges.",
        "المبالغ أعلاه حسب الرسوم الفعلية المعلنة من الهيئة العامة للقوى العاملة ووزارة الداخلية ووزارة الصحة "
        "ووزارة الخارجية والجهات الأخرى، وهي قابلة للتغيير في حال تعديل هذه الجهات لرسومها.")
# الجهة (من جدول الرسوم) بالإنجليزي
AUTHORITY_EN = {
    "الهيئة العامة للقوى العاملة": "Public Authority for Manpower (PAM)",
    "وزارة الداخلية": "Ministry of Interior",
    "وزارة الصحة": "Ministry of Health",
    "وزارة الخارجية": "Ministry of Foreign Affairs",
    "الهيئة العامة للمعلومات المدنية": "Public Authority for Civil Information (PACI)",
}


# ---------------------------------------------------------------------------
# اللغة
# ---------------------------------------------------------------------------
def _lang(ws):
    return getattr(ws, "_lunx_lang", "both")


def _set_lang(ws, lang):
    ws._lunx_lang = lang if lang in custody.DOC_LANGS else "both"
    ws.sheet_view.rightToLeft = ws._lunx_lang == "ar"


def _L(ws, en, ar, sep=" / "):
    """نص ثابت بلغة الشيت: إنجليزي، عربي، أو الاتنين (بـ sep)."""
    lang = _lang(ws)
    return ar if lang == "ar" else en if lang == "en" else f"{en}{sep}{ar}"


def _V(ws, en, ar, sep="  —  "):
    """قيمة (اسم) بلغة الشيت — اللي ناقص بياخد التاني، والثنائي بيكتب الاتنين لو مختلفين."""
    en, ar = en or ar or "", ar or en or ""
    lang = _lang(ws)
    return ar if lang == "ar" else en if lang == "en" else (en if en == ar else f"{en}{sep}{ar}")


def _count(ws, n, en, ar):
    """«3 Employees · موظف»"""
    return _L(ws, f"{n} {en}", f"{n} {ar}", " · ").replace(f" · {n} ", " · ")


# ---------------------------------------------------------------------------
# أدوات التنسيق
# ---------------------------------------------------------------------------
def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _side(color=LINE, style="thin"):
    return Side(style=style, color=color)


GRID = Border(left=_side(), right=_side(), top=_side(), bottom=_side())


def _put(ws, row, col, value=None, *, to_col=None, to_row=None, bold=False, size=10, color=INK, fill=None, h="center",
         v="center", wrap=True, fmt=None, border=None, italic=False):
    """قيمة في خانة (ومدمجة لحد to_col / to_row) بتنسيقها كله في سطر واحد. الشيت العربي (من اليمين للشمال)
    بيعكس المحاذاة: «بداية السطر» يمين."""
    if _lang(ws) == "ar" and h in ("left", "right"):
        h = "right" if h == "left" else "left"
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, readingOrder=2 if _lang(ws) == "ar" else 0)
    if fill:
        c.fill = _fill(fill)
    if fmt:
        c.number_format = fmt
    last_col, last_row = to_col or col, to_row or row
    if (last_col, last_row) != (col, row):
        ws.merge_cells(start_row=row, start_column=col, end_row=last_row, end_column=last_col)
    if border or fill:
        for r in range(row, last_row + 1):
            for cc in range(col, last_col + 1):
                x = ws.cell(row=r, column=cc)
                if border:
                    x.border = border
                if fill:
                    x.fill = _fill(fill)
    return c


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _sheet_name(wb, name):
    name = re.sub(r"[\[\]:*?/\\]", "-", (name or "—").strip())[:28] or "—"
    base, n = name, 2
    while name in wb.sheetnames:
        name = f"{base[:25]} ({n})"
        n += 1
    return name


def _logo(ws, company, anchor="A1", height=58):
    path = db.resolve_file(company.logoPath) if company is not None and company.logoPath else None
    if not (XLImage and path):
        return
    try:
        img = XLImage(path)
    except Exception:                          # صيغة مش مدعومة (SVG …) ← من غير شعار
        return
    ratio = height / float(img.height or height)
    img.height, img.width = height, int((img.width or height) * ratio)
    ws.add_image(img, anchor)


def _header(ws, ncols, company, title, sub=FORM_NO):
    """الشعار · اسم الشركة · عنوان الكشف ورقمه، وتحتهم شريط ذهبي. title / sub = (إنجليزي، عربي) أو نص."""
    for r, hgt in ((1, 24), (2, 22), (3, 18), (4, 4), (5, 8)):
        ws.row_dimensions[r].height = hgt
    _logo(ws, company)
    title_from = ncols - 2 if ncols >= 7 else ncols - 1
    mid_from, mid_to = 3, max(3, title_from - 1)
    name_en = (company.nameEn or company.nameAr) if company is not None else ""
    name_ar = (company.nameAr or company.nameEn) if company is not None else ""
    room = sum(ws.column_dimensions[get_column_letter(c)].width or 10 for c in range(mid_from, mid_to + 1))
    fit = lambda text, size, per_char: max(9, min(size, int(size * room / max(len(text or "") * per_char, 1))))   # اسم طويل ← خط أصغر
    unit = ("Admin Unit", "إدارة الشؤون الإدارية")
    t_en, t_ar = title if isinstance(title, tuple) else (title, title)
    sub = _L(ws, *sub) if isinstance(sub, tuple) else sub
    if _lang(ws) == "both":
        _put(ws, 1, mid_from, name_en, to_col=mid_to, bold=True, size=fit(name_en, 13, 1.2), color=PRIMARY, wrap=False)
        _put(ws, 2, mid_from, name_ar, to_col=mid_to, bold=True, size=fit(name_ar, 12, 0.95), wrap=False)
        _put(ws, 3, mid_from, f"{unit[0]} · {unit[1]}", to_col=mid_to, size=9, color=MUTED)
        _put(ws, 1, title_from, t_en, to_col=ncols, bold=True, size=12, color=PRIMARY, h="right", wrap=False)
        _put(ws, 2, title_from, t_ar, to_col=ncols, bold=True, size=12, h="right", wrap=False)
        _put(ws, 3, title_from, sub, to_col=ncols, size=9, color=MUTED, h="right")
    else:
        ar = _lang(ws) == "ar"
        name = name_ar if ar else name_en
        _put(ws, 1, mid_from, name, to_col=mid_to, to_row=2, bold=True, size=fit(name, 14, 0.95 if ar else 1.2), color=PRIMARY,
             wrap=False)
        _put(ws, 3, mid_from, unit[1] if ar else unit[0], to_col=mid_to, size=9, color=MUTED)
        _put(ws, 1, title_from, t_ar if ar else t_en, to_col=ncols, to_row=2, bold=True, size=15, color=PRIMARY, h="right",
             wrap=False)
        _put(ws, 3, title_from, sub, to_col=ncols, size=9, color=MUTED, h="right")
    for c in range(1, ncols + 1):
        ws.cell(row=4, column=c).fill = _fill(ACCENT)
    return 6


def _info(ws, row, ncols, pairs, box_title, box_note):
    """بيانات الكشف على الشمال (عنوان | قيمة)، وصندوق الإجمالي على اليمين (آخر عمودين). العناوين (إنجليزي، عربي).
    بيرجّع (أول سطر بعدها، خانة رقم الصندوق) — الرقم بيتكتب بعد الجدول (معادلة لخانة الإجمالي)."""
    width = max(sum(ws.column_dimensions[get_column_letter(c)].width or 10 for c in range(3, ncols - 1)) - 1.5, 4)
    for i, (label, value) in enumerate(pairs):
        r = row + i
        ws.row_dimensions[r].height = max(20, 13.5 * math.ceil(len(str(value or "")) * 1.08 / width) + 5)
        _put(ws, r, 1, _L(ws, *label), to_col=2, bold=True, size=9, color=PRIMARY, fill=PRIMARY_SOFT, h="left", border=GRID)
        _put(ws, r, 3, value, to_col=ncols - 2, size=10, h="left", border=GRID)
    last = row + len(pairs) - 1
    ws.row_dimensions[row].height = max(30, ws.row_dimensions[row].height)
    _put(ws, row, ncols - 1, _L(ws, *box_title, sep="\n"), to_col=ncols, bold=True, size=9, color="FFFFFF", fill=PRIMARY,
         border=GRID)
    _put(ws, row + 1, ncols - 1, None, to_col=ncols, to_row=max(row + 1, last - 1), bold=True, size=18, color=PRIMARY,
         fmt=KWD, border=GRID)
    _put(ws, last, ncols - 1, box_note, to_col=ncols, size=9, color=MUTED, fill=PRIMARY_SOFT, border=GRID)
    return last + 2, ws.cell(row=row + 1, column=ncols - 1)


def _table_head(ws, row, headers):
    """headers = نصوص أو (إنجليزي، عربي)."""
    ws.row_dimensions[row].height = 46 if _lang(ws) == "both" else 32
    for i, text in enumerate(headers, 1):
        _put(ws, row, i, _L(ws, *text, sep="\n") if isinstance(text, tuple) else text, bold=True, size=9, color="FFFFFF",
             fill=PRIMARY, border=Border(left=_side("FFFFFF"), right=_side("FFFFFF"), top=_side(PRIMARY), bottom=_side(PRIMARY)))


def _lines(ws, col, text):
    if not isinstance(text, str) or text.startswith("="):
        return 1
    width = max((ws.column_dimensions[get_column_letter(col)].width or 10) - 1.5, 4)
    return sum(max(1, math.ceil(len(part) * 1.08 / width)) for part in text.split("\n"))


def _table_row(ws, row, values, money_from, zebra, left_cols=(2,)):
    ws.row_dimensions[row].height = max(20, 13.5 * max(_lines(ws, i, v) for i, v in enumerate(values, 1)) + 5)
    for i, val in enumerate(values, 1):
        money = i >= money_from and not (isinstance(val, str) and not val.startswith("="))
        _put(ws, row, i, val, size=10, fill=ZEBRA if zebra else None, h="left" if i in left_cols else "center",
             fmt=KWD if money else None, border=GRID, color=MUTED if i == 1 else INK)


def _total_row(ws, row, ncols, label, first_row, last_row, money_from, label_to):
    ws.row_dimensions[row].height = 24
    top = Border(left=_side(), right=_side(), top=_side(PRIMARY, "medium"), bottom=_side(PRIMARY, "double"))
    _put(ws, row, 1, label, to_col=label_to, bold=True, size=10, color=PRIMARY, fill=PRIMARY_SOFT, h="left", border=top)
    for c in range(label_to + 1, ncols + 1):
        col = get_column_letter(c)
        value = f"=SUM({col}{first_row}:{col}{last_row})" if c >= money_from else None
        _put(ws, row, c, value, bold=True, size=10, color=PRIMARY, fill=PRIMARY_SOFT, fmt=KWD, border=top)
    return f"{get_column_letter(ncols)}{row}"


# ---------------------------------------------------------------------------
# المبلغ كتابةً
# ---------------------------------------------------------------------------
def _words(n):
    ones = ("zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen "
            "seventeen eighteen nineteen").split()
    tens = "_ _ twenty thirty forty fifty sixty seventy eighty ninety".split()
    if n < 20:
        return ones[n]
    if n < 100:
        return tens[n // 10] + ("-" + ones[n % 10] if n % 10 else "")
    if n < 1000:
        return ones[n // 100] + " hundred" + (" and " + _words(n % 100) if n % 100 else "")
    for div, name in ((10 ** 9, "billion"), (10 ** 6, "million"), (1000, "thousand")):
        if n >= div:
            rest = n % div
            return _words(n // div) + " " + name + ((" and " if rest < 100 else " ") + _words(rest) if rest else "")
    return str(n)


def amount_words(amount):
    """170.25 ← «One hundred and seventy Kuwaiti Dinars and 250 Fils only»"""
    kd, fils = divmod(round(float(amount or 0) * 1000), 1000)
    text = f"{_words(kd).capitalize()} Kuwaiti Dinar{'' if kd == 1 else 's'}"
    return text + (f" and {fils} Fils" if fils else "") + " only"


AR_ONES = ["", "واحد", "اثنان", "ثلاثة", "أربعة", "خمسة", "ستة", "سبعة", "ثمانية", "تسعة", "عشرة", "أحد عشر", "اثنا عشر",
           "ثلاثة عشر", "أربعة عشر", "خمسة عشر", "ستة عشر", "سبعة عشر", "ثمانية عشر", "تسعة عشر"]
AR_TENS = ["", "", "عشرون", "ثلاثون", "أربعون", "خمسون", "ستون", "سبعون", "ثمانون", "تسعون"]
AR_HUNDREDS = ["", "مائة", "مائتان", "ثلاثمائة", "أربعمائة", "خمسمائة", "ستمائة", "سبعمائة", "ثمانمائة", "تسعمائة"]


def _ar_below_1000(n):
    h, r = divmod(n, 100)
    parts = [AR_HUNDREDS[h]] if h else []
    if r:
        parts.append(AR_ONES[r] if r < 20 else (AR_ONES[r % 10] + " و" if r % 10 else "") + AR_TENS[r // 10])
    return " و".join(parts)


def _ar_words(n):
    """275 ← «مائتان وخمسة وسبعون»، 3500 ← «ثلاثة آلاف وخمسمائة»"""
    if n == 0:
        return "صفر"
    out = []
    for div, one, two, plural in ((10 ** 6, "مليون", "مليونان", "ملايين"), (1000, "ألف", "ألفان", "آلاف")):
        q, n = divmod(n, div)
        if q == 1:
            out.append(one)
        elif q == 2:
            out.append(two)
        elif 3 <= q <= 10:
            out.append(f"{_ar_below_1000(q)} {plural}")
        elif q > 10:
            out.append(f"{_ar_below_1000(q)} {one}")
    if n:
        out.append(_ar_below_1000(n))
    return " و".join(out)


def amount_words_ar(amount):
    """170.25 ← «فقط مائة وسبعون دينار كويتي ومائتان وخمسون فلس لا غير»"""
    kd, fils = divmod(round(float(amount or 0) * 1000), 1000)
    return f"فقط {_ar_words(kd)} دينار كويتي" + (f" و{_ar_words(fils)} فلس" if fils else "") + " لا غير"


def _words_row(ws, row, ncols, amount):
    lang = _lang(ws)
    en, ar = f"Amount in words:  {amount_words(amount)}", f"المبلغ كتابةً:  {amount_words_ar(amount)}"
    ws.row_dimensions[row].height = 32 if lang == "both" else 20
    _put(ws, row, 1, ar if lang == "ar" else en if lang == "en" else f"{en}\n{ar}", to_col=ncols, size=9, italic=True,
         color=INK, h="left", border=Border(bottom=_side()))
    return row + 1


def _note(ws, row, ncols, text=NOTE):
    ws.row_dimensions[row].height = 36 if _lang(ws) == "both" else 26
    _put(ws, row, 1, _L(ws, *text, sep="\n"), to_col=ncols, size=8, italic=True, color=MUTED, h="left", v="top")
    return row + 2


def _spans(ws, ncols, n):
    """n خانة على عرض الكشف بالتساوي تقريبًا (حسب عرض الأعمدة، مش عددها)."""
    widths = [ws.column_dimensions[get_column_letter(c)].width or 10 for c in range(1, ncols + 1)]
    target, spans, start, acc = sum(widths) / n, [], 1, 0
    for c in range(1, ncols + 1):
        acc += widths[c - 1]
        left_boxes = n - len(spans) - 1
        if left_boxes and ncols - c >= left_boxes and acc >= target * (len(spans) + 1) - widths[c - 1] / 2:
            spans.append((start, c))
            start = c + 1
    spans.append((start, ncols))
    return spans


def _signatures(ws, row, ncols, labels, prepared_by=None):
    """جدول توقيعات: العنوان، مساحة التوقيع، الاسم، التاريخ — كل خانة بإطار. labels = [(إنجليزي، عربي)]."""
    for r, hgt in ((row, 22), (row + 1, 42), (row + 2, 18), (row + 3, 18)):
        ws.row_dimensions[r].height = hgt
    for i, ((a, b), label) in enumerate(zip(_spans(ws, ncols, len(labels)), labels)):
        _put(ws, row, a, _L(ws, *label), to_col=b, bold=True, size=9, color=PRIMARY, fill=PRIMARY_SOFT, border=GRID)
        _put(ws, row + 1, a, None, to_col=b, border=GRID)
        name = prepared_by if (i == 0 and prepared_by) else "………………………"
        _put(ws, row + 2, a, f"{_L(ws, 'Name', 'الاسم')}:  {name}", to_col=b, size=8, color=MUTED, h="left", border=GRID)
        _put(ws, row + 3, a, f"{_L(ws, 'Date', 'التاريخ')}:  ……/……/……", to_col=b, size=8, color=MUTED, h="left", border=GRID)
    return row + 4


def _print_setup(ws, landscape, title_row, footer_left, user=None):
    """A4 على عرض الصفحة، رأس الجدول بيتكرر، والتذييل: رقم الكشف، «صفحة X من Y»، وختم «صادر بواسطة المستخدم — الوقت» (من غير اسم البرنامج)."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top, ws.page_margins.bottom = 0.45, 0.55
    if title_row:
        ws.print_title_rows = f"{title_row}:{title_row}"
    who = (user or "").replace("&", "&&")
    ws.oddFooter.left.text, ws.oddFooter.left.size = footer_left.replace("&", "&&"), 8
    ws.oddFooter.center.text, ws.oddFooter.center.size = _L(ws, "Page &P of &N", "صفحة &P من &N", " · "), 8
    ws.oddFooter.right.text, ws.oddFooter.right.size = _L(ws, f"Issued by {who} · &D &T", f"صادر بواسطة {who} · &D &T",
                                                          " / ") if _lang(ws) != "both" else f"{who} · &D &T", 8      # من غير اسم البرنامج
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PRIMARY


# مساحة الطباعة (نقطة) على A4 بالهوامش اللي فوق: (العرض، الطول)
PRINTABLE = {False: (544.9, 769.9), True: (791.5, 523.3)}


def _fit(ws, landscape, title_row, tail_row, one_page=1.15):
    """لو الشيت أطول من صفحة بحبة (لحد one_page) ← يتصغّر ويبقى صفحة واحدة. غير كده: لو الجزء الأخير (من tail_row: المبلغ كتابةً
    والملاحظة والتوقيعات) هيتقسم على صفحتين، بيبدأ صفحة جديدة. الحساب تقريبي بنفس طريقة إكسيل (عرض الأعمدة ← التصغير)."""
    from openpyxl.worksheet.pagebreak import Break
    pw, ph = PRINTABLE[landscape]
    width = sum(((ws.column_dimensions[get_column_letter(c)].width or 8.43) * 7 + 5) * 0.75 for c in range(1, ws.max_column + 1))
    scale = min(1.0, pw / width)
    heights = [(ws.row_dimensions[r].height or 15) * scale for r in range(1, ws.max_row + 1)]
    if sum(heights) <= ph * one_page:
        ws.page_setup.fitToHeight = 1
        return
    repeat = heights[title_row - 1] if title_row else 0
    used = 0
    for r in range(1, tail_row):
        h = heights[r - 1]
        used = repeat + h if used + h > ph else used + h
    block = sum(heights[tail_row - 1:])
    if used + block > ph and block <= ph - repeat:
        ws.row_breaks.append(Break(id=tail_row - 1))


# ---------------------------------------------------------------------------
# البيانات
# ---------------------------------------------------------------------------
def _fmt(d):
    d = db.parse_date(d) if isinstance(d, str) else d
    return d.strftime("%d/%m/%Y") if d else ""


def _person(s, kind, pid, snap_name, tr):
    """الاسم بالإنجليزي والعربي والمهنة (عربي وإنجليزي) والإقامة من السجل الحالي (والاسم وقت الطلب لو اتمسح)."""
    rec = s.get(M.Employee, pid) if kind == "employee" else s.get(M.Candidate, pid)
    prof = rec.profession if rec is not None else ""
    prof_en = (getattr(rec, "professionEn", None) if rec is not None else None) or value_i18n.lookup(tr, "profession", prof)
    return {"name": (rec.nameEn if rec is not None and rec.nameEn else None) or snap_name or pid,
            "nameAr": (rec.name if rec is not None else None) or snap_name or pid, "profession": prof or "",
            "professionEn": prof_en or prof or "",
            "residencyExp": getattr(rec, "residencyExp", None) if rec is not None else None}


def _pname(ws, p):
    """اسم الشخص في الجداول: العربي في الكشف العربي، والإنجليزي في الإنجليزي والثنائي."""
    return p["nameAr"] if _lang(ws) == "ar" else p["name"]


def _pprof(ws, p):
    return p["professionEn"] if _lang(ws) == "en" else p["profession"]


def _cost_center(s, name, fallback_company):
    """مركز التكلفة ← (اسمه بالإنجليزي، بالعربي، شركته) — الشركة للشعار والاسم في رأس الكشف."""
    cc = s.scalar(select(M.CostCenter).where(M.CostCenter.name == name)) if name else None
    co_id = (cc.companyId if cc is not None else None) or fallback_company
    en, ar = ((cc.nameEn if cc is not None else None) or name, name) if name else ("No cost center", "بدون مركز تكلفة")
    return en, ar, (s.get(M.Company, co_id) if co_id else None)


def _items(lines):
    cols = []
    for ln in sorted(lines, key=lambda x: x.position):
        key = ln.feeItemId or ln.itemName
        if key not in [c["key"] for c in cols]:
            cols.append({"key": key, "en": ln.itemNameEn or ln.itemName, "ar": ln.itemName})
    return cols


def _group_people(s, lines):
    """البنود ← [(مركز التكلفة، [أشخاص])] — الشخص: بياناته وبنوده."""
    tr, people = value_i18n.merged(s), {}
    for ln in lines:
        k = (ln.personKind, ln.personId)
        if k not in people:
            people[k] = {**_person(s, ln.personKind, ln.personId, ln.personName, tr), "personId": ln.personId, "civilId": ln.civilId or "",
                         "costCenter": ln.costCenter, "companyId": ln.companyId, "lines": []}
        people[k]["lines"].append(ln)
    groups = {}
    for p in people.values():
        groups.setdefault(p["costCenter"] or "", []).append(p)
    return sorted(((cc, sorted(ps, key=lambda p: p["name"].lower())) for cc, ps in groups.items()), key=lambda g: g[0])


def _amount(ln, actual=True):
    return float((ln.actual if actual and ln.actual is not None else ln.planned) or 0)


def _tx(ws, c):
    tx = custody.TX_TYPES.get(c.txType, {})
    return _V(ws, tx.get("en", c.txType), tx.get("label", c.txType))


def _company(ws, co):
    return _V(ws, co.nameEn, co.nameAr) if co is not None else "—"


# ---------------------------------------------------------------------------
# الفاتورة (لمركز تكلفة في تقفيل)
# ---------------------------------------------------------------------------
def _issuer(s):
    """الشركة المُصدِرة (أبراج انرجي — «إعدادات الفواتير»): شعارها واسمها على الطلب والفواتير والملخص."""
    co = custody.settings(s)["issuerCompanyId"]
    return s.get(M.Company, co) if co else None


def _status(ws, inv):
    if inv.status == "approved":
        return f"{_L(ws, 'Approved by Accounts', 'اعتمدتها الحسابات', ' · ')} — {_fmt(inv.approvedDate)}"
    return _L(ws, "Pending Accounts Approval", "بانتظار اعتماد الحسابات", " · ")


def _invoice_note(support):
    en = "Government charges paid on behalf of the above cost center as per the attached original receipts"
    ar = "رسوم حكومية مدفوعة نيابةً عن مركز التكلفة أعلاه حسب الإيصالات الأصلية المرفقة"
    if support:
        en, ar = en + ", plus administrative support per employee", ar + "، بالإضافة إلى الدعم الإداري لكل موظف"
    return f"{en}.", f"{ar}."


def _closed_by(inv):
    """«a.ahmed — 29/09/2026 11:53» (مين قفّل وإمتى)"""
    return f"{inv.createdBy or '—'} — {inv.createdAt.strftime('%d/%m/%Y %H:%M')}" if inv.createdAt else inv.createdBy or "—"


def _invoice_lines(inv, lines):
    return [ln for ln in lines if (ln.costCenter or "") == (inv.costCenter or "")]


def _invoice_sheet(ws, s, c, inv, lines, issuer, user):
    """فاتورة من المُصدِر لمركز تكلفة: موظف في كل صف وبند في كل عمود (+ الدعم الإداري لو المركز عليه دعم)،
    وملخص الفاتورة (الرسوم الحكومية + الدعم = إجمالي المستحق) والمبلغ كتابةً وتوقيعات الاعتماد."""
    people = [p for _, ps in _group_people(s, lines) for p in ps]
    items = _items(lines)
    # الدعم الإداري مرة لكل موظف في العهدة: في أول تقفيل فيه إجراء ليه — اللي اتحسب له قبل كده «—»
    earlier = custody.earlier_closed(s, c, inv.closingDate)
    fee = float(inv.supportFee or 0)
    charged = [p for p in people if p["personId"] not in earlier]
    support = fee if fee and charged else 0.0
    headers = ["#", ("Employee Name", "اسم الموظف"), ("Profession", "المهنة"), ("Civil ID", "الرقم المدني")] \
        + [(it["en"], it["ar"]) for it in items] + ([("Admin Support", "الدعم الإداري")] if support else []) \
        + [("Total (KWD)", "الإجمالي (د.ك)")]
    n = len(headers)
    last_item = 4 + len(items)
    _widths(ws, [5, 34, 22, 15] + [15] * len(items) + ([14] if support else []) + [15])
    number = custody.invoice_no(inv)
    row = _header(ws, n, issuer, ("INVOICE", "فاتورة"), number)
    cc_en, cc_ar, bill_co = _cost_center(s, inv.costCenter, inv.billCompanyId)
    cc_label = f"{inv.ccCode} — {_V(ws, cc_en, cc_ar)}" if inv.ccCode else _V(ws, cc_en, cc_ar)
    row, box = _info(ws, row, n, [
        (("Invoice No.", "رقم الفاتورة"), number),
        (("Invoice Date", "تاريخ الفاتورة"), _fmt(inv.closingDate)),
        (("Bill To", "فاتورة إلى"), cc_label),
        (("Company", "الشركة"), _company(ws, bill_co)),
        (("Request For", "نوع المعاملة"), _tx(ws, c)),
        (("Custody / Closing", "العهدة / التقفيل"), f"{custody.custody_no(c)}   ·   {inv.closingRef or '—'}"),
        (("Closed By", "أُقفلت بواسطة"), _closed_by(inv)),
        (("Status", "الحالة"), _status(ws, inv)),
    ], ("TOTAL DUE (KWD)", "إجمالي المستحق"), _count(ws, len(people), "Employees", "موظف"))
    ws.cell(row=row - 2, column=3).font = Font(name=FONT, size=10, bold=True,
                                              color=PRIMARY if inv.status == "approved" else ACCENT)
    head = row
    _table_head(ws, head, headers)
    first = head + 1
    for i, p in enumerate(people):
        by = {ln.feeItemId or ln.itemName: ln for ln in p["lines"]}
        r = first + i
        amounts = [(_amount(by[it["key"]]) if it["key"] in by else "—") for it in items]
        _table_row(ws, r, [i + 1, _pname(ws, p), _pprof(ws, p), p["civilId"], *amounts,
                           *([support if p["personId"] not in earlier else "—"] if support else []),
                           f"=SUM({get_column_letter(5)}{r}:{get_column_letter(n - 1)}{r})"], money_from=5, zebra=i % 2 == 1)
    last = first + len(people) - 1
    tot = last + 1
    _total_row(ws, tot, n, f"{_L(ws, 'TOTAL', 'الإجمالي')}  ({len(people)})", first, last, 5, 4)
    # ملخص الفاتورة على الجنب
    L = get_column_letter
    label_from = max(2, n - 3)
    parts = [(_L(ws, "Government Charges", "الرسوم الحكومية"), f"=SUM(E{tot}:{L(last_item)}{tot})")]
    if support:
        parts.append((f"{_L(ws, 'Admin Support', 'الدعم الإداري')}  ({len(charged)} × {support:,.3f})", f"={L(n - 1)}{tot}"))
    r = tot + 2
    for i, (label, value) in enumerate(parts):
        ws.row_dimensions[r + i].height = 20
        _put(ws, r + i, label_from, label, to_col=n - 1, size=10, h="left", border=GRID)
        _put(ws, r + i, n, value, size=10, fmt=KWD, border=GRID)
    due = r + len(parts)
    ws.row_dimensions[due].height = 26
    _put(ws, due, label_from, _L(ws, "TOTAL DUE", "إجمالي المستحق"), to_col=n - 1, bold=True, size=11, color="FFFFFF",
         fill=PRIMARY, h="left", border=GRID)
    _put(ws, due, n, f"=SUM({L(n)}{r}:{L(n)}{due - 1})", bold=True, size=11, color="FFFFFF", fill=PRIMARY, fmt=KWD, border=GRID)
    box.value = f"={L(n)}{due}"
    total = sum(_amount(ln) for ln in lines) + support * len(charged)
    row = _words_row(ws, due + 2, n, total)
    row = _note(ws, row + 1, n, _invoice_note(support))
    _signatures(ws, row + 1, n, [("Prepared By", "أعده"), ("Manager Approval", "اعتماد المدير"),
                                 ("Accounts Approval", "اعتماد الحسابات")], inv.createdBy or user)
    ws.freeze_panes = ws.cell(row=first, column=3)
    _print_setup(ws, n > 7, head, f"{number} · {cc_label}", user)
    _fit(ws, n > 7, head, tot + 2, one_page=1.4)                   # الفاتورة في صفحة واحدة لو ينفع
    return people


def _details_sheet(wb, s, c, inv, lines, issuer, user, lang):
    """ملحق الفاتورة: كشف فردي لكل موظف (الإجراء، الجهة، المبلغ، رقم الإيصال) — كل كشف في صفحة لوحده."""
    from openpyxl.worksheet.pagebreak import Break
    number = custody.invoice_no(inv)
    ws = wb.create_sheet(_sheet_name(wb, f"{number[4:]} {'تفاصيل' if lang == 'ar' else 'Details'}"))
    _set_lang(ws, lang)
    cc_en, cc_ar, _ = _cost_center(s, inv.costCenter, inv.billCompanyId)
    cc_label = f"{inv.ccCode} — {_V(ws, cc_en, cc_ar)}" if inv.ccCode else _V(ws, cc_en, cc_ar)
    r = 1
    earlier = custody.earlier_closed(s, c, inv.closingDate)       # الدعم الإداري اتحسب لهم في تقفيل قبل كده
    for i, p in enumerate(p for _, ps in _group_people(s, lines) for p in ps):
        if i:
            ws.row_breaks.append(Break(id=r - 1))
        fee = 0.0 if p["personId"] in earlier else float(inv.supportFee or 0)
        r, _, _ = _closing_individual(ws, r, c, cc_label, issuer, p, fee, inv.closingDate,
                                      inv.createdBy or user, i == 0, number, inv.closingRef)
    _print_setup(ws, False, None, f"{number} · {cc_label}", user)
    ws.sheet_properties.tabColor = ACCENT


def _closing_individual(ws, row, c, cc_label, company, p, admin_fee, when, user, first_page, number, c_ref=""):
    """كشف فردي لموظف (نموذج T-104) — بيبدأ من row، وكل كشف في صفحة مطبوعة لوحده."""
    n = 6
    c_ref = c_ref or custody.custody_no(c)
    if first_page:
        _widths(ws, [5, 46, 28, 16, 16, 20])
    title = ("EMPLOYEE STATEMENT", "كشف موظف")
    row = _header(ws, n, company, title) if first_page else _header_at(ws, row, n, company, title)
    lines = sorted(p["lines"], key=lambda x: x.position)
    row, box = _info(ws, row, n, [
        (("Employee Name", "اسم الموظف"), _V(ws, p["name"], p["nameAr"])),
        (("Profession", "المهنة"), _V(ws, p["professionEn"], p["profession"])),
        (("Civil ID", "الرقم المدني"), p["civilId"]),
        (("Cost Center", "مركز التكلفة"), cc_label),
        (("Request For", "نوع المعاملة"), _tx(ws, c)),
        (("Invoice No.", "رقم الفاتورة"), f"{number}   ·   {c_ref}"),
        (("Invoice Date", "تاريخ الفاتورة"), _fmt(when)),
    ], ("TOTAL (KWD)", "الإجمالي"), _L(ws, "Kuwaiti Dinar", "دينار كويتي", " · "))
    head = row
    _table_head(ws, head, ["#", ("Process", "الإجراء"), ("Authority", "الجهة"), ("Charges (KWD)", "المبلغ (د.ك)"),
                           ("Receipt No.", "رقم الإيصال"), ("Notes", "ملاحظات")])
    for i, ln in enumerate(lines):
        authority = AUTHORITY_EN.get(ln.authority or "", ln.authority or "") if _lang(ws) == "en" else ln.authority or ""
        _table_row(ws, head + 1 + i, [i + 1, _V(ws, ln.itemNameEn, ln.itemName), authority, _amount(ln), ln.receiptNo or "", ""],
                   money_from=4, zebra=i % 2 == 1, left_cols=(2, 3))
    r = head + len(lines)
    if admin_fee:
        r += 1
        _table_row(ws, r, [len(lines) + 1, _V(ws, "Administrative Support", "الدعم الإداري"),
                           _L(ws, "Admin Unit", "الشؤون الإدارية", " · "), admin_fee, "", ""],
                   money_from=4, zebra=len(lines) % 2 == 1, left_cols=(2, 3))
        for cc in (5, 6):                                     # مابيتجمعوش
            ws.cell(row=r, column=cc).number_format = "General"
    tot = r + 1
    top = Border(left=_side(), right=_side(), top=_side(PRIMARY, "medium"), bottom=_side(PRIMARY, "double"))
    ws.row_dimensions[tot].height = 24
    _put(ws, tot, 1, _L(ws, "TOTAL", "الإجمالي"), to_col=3, bold=True, color=PRIMARY, fill=PRIMARY_SOFT, h="left", border=top)
    _put(ws, tot, 4, f"=SUM(D{head + 1}:D{r})", bold=True, color=PRIMARY, fill=PRIMARY_SOFT, fmt=KWD, border=top)
    box.value = f"=D{tot}"
    _put(ws, tot, 5, None, to_col=6, fill=PRIMARY_SOFT, border=top)
    grand = sum(_amount(ln) for ln in lines) + admin_fee
    row = _words_row(ws, tot + 2, n, grand)
    row = _note(ws, row + 1, n)
    row = _signatures(ws, row, n, [("Prepared By", "أعده"), ("Manager Approval", "اعتماد المدير"),
                                   ("Accounts Approval", "اعتماد الحسابات")], user)
    return row + 1, head, grand


def _header_at(ws, row, n, company, title):
    """رأس كشف فردي تاني في نفس الشيت (بعد فاصل صفحة): نفس الرأس من غير الشعار."""
    ws.row_dimensions[row].height = 24
    _put(ws, row, 1, _company(ws, company) if _lang(ws) != "both" else (company.nameEn if company is not None else ""),
         to_col=3, bold=True, size=12, color=PRIMARY, h="left")
    _put(ws, row, 4, _L(ws, *title, sep=" · "), to_col=n, bold=True, size=11, color=PRIMARY, h="right")
    if _lang(ws) == "both":
        _put(ws, row + 1, 1, company.nameAr if company is not None else "", to_col=3, size=10, h="left")
    _put(ws, row + 1, 4, _L(ws, *FORM_NO), to_col=n, size=9, color=MUTED, h="right")
    for c in range(1, n + 1):
        ws.cell(row=row + 2, column=c).fill = _fill(ACCENT)
    ws.row_dimensions[row + 2].height = 4
    return row + 4


def _closed(s, c, when):
    return s.scalars(select(M.CustodyLine).where(M.CustodyLine.custodyId == c.id, M.CustodyLine.closedDate == when)).all()


def _save(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def closing_workbook(s, c, when, user, individual=False, lang="both"):
    """تقفيل بتاريخه: «الملخص» (قايمة الفواتير + تسوية العهدة مع المستلم) وشيت فاتورة لكل مركز تكلفة
    (+ ملحق كشف فردي لكل موظف لو individual)."""
    invoices = s.scalars(select(M.Invoice).where(M.Invoice.custodyId == c.id, M.Invoice.closingDate == when)
                         .order_by(M.Invoice.no)).all()
    lines = _closed(s, c, when)
    issuer = _issuer(s)
    wb = Workbook()
    summary = wb.active
    _set_lang(summary, lang)
    rows = []
    for inv in invoices:
        inv_lines = _invoice_lines(inv, lines)
        ws = wb.create_sheet(_sheet_name(wb, f"{custody.invoice_no(inv)[4:]} {inv.costCenter or ''}"))
        _set_lang(ws, lang)
        cc_en, cc_ar, bill_co = _cost_center(s, inv.costCenter, inv.billCompanyId)
        _invoice_sheet(ws, s, c, inv, inv_lines, issuer, user)
        if individual:
            _details_sheet(wb, s, c, inv, inv_lines, issuer, user, lang)
        rows.append((inv, f"{inv.ccCode} — {_V(ws, cc_en, cc_ar)}" if inv.ccCode else _V(ws, cc_en, cc_ar), bill_co, ws.title))
    _closing_summary(summary, s, c, rows, issuer, when, user, invoices[0] if invoices else None)
    return _save(wb)


def invoice_workbook(s, inv, user, individual=False, lang="both"):
    """فاتورة واحدة (+ الملحق الفردي)."""
    c = s.get(M.Custody, inv.custodyId)
    lines = _invoice_lines(inv, _closed(s, c, inv.closingDate))
    issuer = _issuer(s)
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(wb, custody.invoice_no(inv))
    _set_lang(ws, lang)
    _invoice_sheet(ws, s, c, inv, lines, issuer, user)
    if individual:
        _details_sheet(wb, s, c, inv, lines, issuer, user, lang)
    return _save(wb)


def _closing_summary(ws, s, c, rows, issuer, when, user, first=None):
    """الملخص الداخلي بين المستلم والشركة المُصدِرة: الفواتير اللي طلعت من التقفيل، وتسوية العهدة كلها."""
    ws.title = _L(ws, "Summary", "الملخص", " ")
    headers = ["#", ("Invoice No.", "رقم الفاتورة"), ("Bill To (Cost Center)", "فاتورة إلى"), ("Company", "الشركة"),
               ("Status", "الحالة"), ("Employees", "الموظفين"), ("Government", "الرسوم الحكومية"),
               ("Admin Support", "الدعم الإداري"), ("Total (KWD)", "الإجمالي (د.ك)")]
    n = len(headers)
    _widths(ws, [5, 20, 26, 28, 13, 11, 14, 13, 15])
    number = custody.custody_no(c)
    row = _header(ws, n, issuer, ("CUSTODY CLOSING SUMMARY", "ملخص تقفيل العهدة"))
    row, box = _info(ws, row, n, [
        (("Custody / Closing", "العهدة / التقفيل"), f"{number}   ·   {first.closingRef if first is not None and first.closingRef else '—'}"),
        (("Request For", "نوع المعاملة"), _tx(ws, c)),
        (("Custodian", "المستلم"), c.custodian),
        (("Closed By", "أُقفلت بواسطة"), _closed_by(first) if first is not None else user),
        (("Disbursed", "تاريخ الصرف"), f"{_fmt(c.disbursedDate)}   ·   {c.disbursedAmount or 0:,.3f} {_L(ws, 'KWD', 'د.ك', ' ')}"),
        (("Closing Date", "تاريخ التقفيل"), _fmt(when)),
    ], ("TOTAL INVOICED (KWD)", "إجمالي الفواتير"), _count(ws, len(rows), "Invoices", "فاتورة"))
    head = row
    _table_head(ws, head, headers)
    for i, (inv, label, company, sheet) in enumerate(rows):
        r = head + 1 + i
        status = f"{_L(ws, 'Approved', 'معتمدة', ' · ')}\n{_fmt(inv.approvedDate)}" if inv.status == "approved" \
            else _L(ws, "Pending", "بانتظار الحسابات", " · ")
        _table_row(ws, r, [i + 1, custody.invoice_no(inv), label, _company(ws, company) if _lang(ws) != "both"
                           else (company.nameEn if company is not None else ""), status,
                           inv.employees, inv.govAmount, inv.supportAmount, f"=G{r}+H{r}"],
                   money_from=6, zebra=i % 2 == 1, left_cols=(3, 4))
        ws.cell(row=r, column=2).hyperlink = f"#'{sheet}'!A1"
        ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True, color=PRIMARY, underline="single")
        ws.cell(row=r, column=5).font = Font(name=FONT, size=9, bold=True, color=PRIMARY if inv.status == "approved" else ACCENT)
        ws.cell(row=r, column=6).number_format = "0"
    last = head + len(rows)
    tot = last + 1
    box.value = "=" + _total_row(ws, tot, n, f"{_L(ws, 'TOTAL', 'الإجمالي')}  ({len(rows)})", head + 1, last, 6, 4)
    ws.cell(row=tot, column=6).number_format = "0"
    # تسوية العهدة مع المستلم (العهدة كلها، مش التقفيل ده بس) — الدعم الإداري على مراكز التكلفة، مش من فلوس العهدة
    lines = s.scalars(select(M.CustodyLine).where(M.CustodyLine.custodyId == c.id)).all()
    spent = sum(_amount(ln) for ln in lines if ln.done)
    got = float(c.disbursedAmount or 0)
    r = tot + 2
    _put(ws, r, 1, _L(ws, "CUSTODY SETTLEMENT", "تسوية العهدة مع المستلم", " · "), to_col=n, bold=True, size=10, color="FFFFFF",
         fill=PRIMARY, h="left")
    for i, (label, val, col) in enumerate(((("Amount disbursed to custodian", "المصروف للمستلم"), got, INK),
                                           (("Government charges executed (whole custody)", "المنفّذ من العهدة"), spent, INK),
                                           (("Balance with custodian", "الرصيد مع المستلم") if got >= spent
                                            else ("Due to custodian", "مستحق للمستلم"), got - spent,
                                            PRIMARY if got >= spent else RED))):
        ws.row_dimensions[r + 1 + i].height = 20
        _put(ws, r + 1 + i, 1, _L(ws, *label), to_col=6, size=10, h="left", border=GRID, bold=i == 2)
        _put(ws, r + 1 + i, 7, abs(val) if i == 2 else val, to_col=9, size=10, bold=i == 2, color=col, fmt=KWD, border=GRID)
    row = _note(ws, r + 5, n)
    _signatures(ws, row, n, [("Prepared By", "أعده"), ("Reviewed By", "راجعه"), ("Manager Approval", "اعتماد المدير"),
                             ("Finance", "الإدارة المالية")], user)
    _print_setup(ws, False, head, f"{number} · {c.custodian}", user)
    _fit(ws, False, head, r)


# ---------------------------------------------------------------------------
# طلب صرف العهدة
# ---------------------------------------------------------------------------
def request_workbook(s, c, user, lang="both"):
    lines = s.scalars(select(M.CustodyLine).where(M.CustodyLine.custodyId == c.id)).all()
    items = _items(lines)
    groups = _group_people(s, lines)
    people = [(cc, p) for cc, ps in groups for p in ps]
    employees = custody.TX_TYPES.get(c.txType, {}).get("kind") == "employee"
    wb = Workbook()
    ws = wb.active
    _set_lang(ws, lang)
    ws.title = _L(ws, "Request", "طلب الصرف", " ")
    headers = ["#", ("Employee Name", "اسم الموظف"), ("Profession", "المهنة"), ("Civil ID", "الرقم المدني"),
               ("Cost Center", "مركز التكلفة")] \
        + ([("Residency Expiry", "انتهاء الإقامة"), ("Days Left", "المتبقي")] if employees else []) \
        + [(it["en"], it["ar"]) for it in items] + [("Total (KWD)", "الإجمالي (د.ك)")]
    n = len(headers)
    money_from = 8 if employees else 6
    _widths(ws, [5, 32, 22, 15, 22] + ([14, 10] if employees else []) + [14] * len(items) + [15])
    number = custody.custody_no(c)
    row = _header(ws, n, _issuer(s), ("ADVANCE PAYMENT REQUEST", "طلب صرف عهدة"))
    row, box = _info(ws, row, n, [
        (("Custody No.", "رقم العهدة"), number),
        (("Request For", "نوع المعاملة"), _tx(ws, c)),
        (("Custodian", "المستلم"), c.custodian),
        (("Request Date", "تاريخ الطلب"), _fmt(c.requestDate)),
        (("Notes", "ملاحظات"), c.notes or ""),
    ], ("AMOUNT REQUIRED (KWD)", "المبلغ المطلوب"), _count(ws, len(people), "Employees", "موظف"))
    head = row
    _table_head(ws, head, headers)
    today = date.today()
    cc_names = {}
    for i, (cc, p) in enumerate(people):
        r = head + 1 + i
        by = {ln.feeItemId or ln.itemName: ln for ln in p["lines"]}
        exp = p["residencyExp"]
        extra = [_fmt(exp), (exp - today).days if exp else ""] if employees else []
        amounts = [(_amount(by[it["key"]], actual=False) if it["key"] in by else "—") for it in items]
        if cc not in cc_names:
            en, ar, _ = _cost_center(s, cc, None)
            cc_names[cc] = (ar if _lang(ws) == "ar" else en) if cc else "—"
        _table_row(ws, r, [i + 1, _pname(ws, p), _pprof(ws, p), p["civilId"], cc_names[cc], *extra, *amounts,
                           f"=SUM({get_column_letter(money_from)}{r}:{get_column_letter(n - 1)}{r})"],
                   money_from=money_from, zebra=i % 2 == 1, left_cols=(2,))
        if employees and exp and (exp - today).days < 0:
            ws.cell(row=r, column=7).font = Font(name=FONT, size=10, bold=True, color=RED)
    last = head + len(people)
    tot = last + 1
    box.value = "=" + _total_row(ws, tot, n, f"{_L(ws, 'GRAND TOTAL', 'الإجمالي العام')}  ({len(people)})", head + 1, last,
                                  money_from, money_from - 1)
    # ملخص المطلوب لكل مركز تكلفة
    r = tot + 2
    _put(ws, r, 1, _L(ws, "SUMMARY PER COST CENTER", "ملخص المطلوب لكل مركز تكلفة", " · "), to_col=n, bold=True, size=10,
         color="FFFFFF", fill=PRIMARY, h="left")
    for i, (cc, ps) in enumerate(groups):
        en, ar, _ = _cost_center(s, cc, ps[0]["companyId"])
        amt = sum(_amount(ln, actual=False) for p in ps for ln in p["lines"])
        _put(ws, r + 1 + i, 1, i + 1, size=10, color=MUTED, border=GRID, fill=ZEBRA if i % 2 else None)
        _put(ws, r + 1 + i, 2, _V(ws, en, ar), to_col=4, size=10, h="left", border=GRID, fill=ZEBRA if i % 2 else None)
        _put(ws, r + 1 + i, 5, _count(ws, len(ps), "Employees", "موظف"), to_col=max(5, n - 1), size=10, border=GRID,
             fill=ZEBRA if i % 2 else None)
        _put(ws, r + 1 + i, n, amt, size=10, fmt=KWD, border=GRID, fill=ZEBRA if i % 2 else None)
    row = _words_row(ws, r + 2 + len(groups), n, sum(_amount(ln, actual=False) for ln in lines))
    row = _note(ws, row + 1, n)
    _signatures(ws, row, n, [("Custodian", "المستلم"), ("Manager", "المسؤول"), ("Finance", "الإدارة المالية")], None)
    ws.freeze_panes = ws.cell(row=head + 1, column=3)
    _print_setup(ws, n > 8, head, f"{number} · {c.custodian}", user)
    _fit(ws, n > 8, head, r)
    return _save(wb)
