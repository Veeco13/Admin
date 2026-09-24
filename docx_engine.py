# -*- coding: utf-8 -*-
"""
محرك عقود العمل (Word)  —  يقابل القسم 11 في وثيقة Lunx
⚠️ مُجمَّد: ممنوع تعديله إلا بطلب صريح.

- resolve_contract_template(emp, company, signatory, ...) → قاموس الحقول (عربي + إنجليزي)
- fill_docx_template(path, context) → bytes لملف docx مُعبّأ
  * يستبدل {{ field }} حتى لو كان مقسومًا على أكثر من run (مع الحفاظ على تنسيق أول run)
  * يدعم MERGEFIELD (يستبدل النص الظاهر للحقل)
- docx_to_html(bytes) → HTML مبسّط للمعاينة
- docx_to_pdf(bytes) → PDF عبر LibreOffice (لو متوفر)
"""
import io
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, date

from docx import Document

PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")

AR_DAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
EN_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# قاموس المهن (عربي ← إنجليزي) — يُستكمل تلقائيًا من بيانات الموظفين
PROFESSION_EN = {
    "مراقب مالي": "Financial Controller",
    "محاسب": "Accountant",
    "سائق / شاحنة": "Driver / Truck",
    "سائق / سيارة خصوصي": "Driver / Private Car",
    "سكرتير": "Secretary",
    "مهندس": "Engineer",
    "فني": "Technician",
    "عامل": "Labourer",
    "مندوب": "Representative",
}

NATIONALITY_EN = {
    "مصر": "Egypt", "مصري": "Egyptian", "الهند": "India", "هندي": "Indian",
    "الأردن": "Jordan", "الاردن": "Jordan", "الجزائر": "Algeria", "الفلبين": "Philippines",
    "باكستان": "Pakistan", "بنغلاديش": "Bangladesh", "نيبال": "Nepal", "كندا": "Canada",
    "لبنان": "Lebanon", "سوريا": "Syria", "ماليزيا": "Malaysia", "سريلانكا": "Sri Lanka",
    "تونس": "Tunisia", "المغرب": "Morocco", "السودان": "Sudan", "اليمن": "Yemen",
    "العراق": "Iraq", "فلسطين": "Palestine", "الكويت": "Kuwait", "السعودية": "Saudi Arabia",
    "بريطانيا": "United Kingdom", "الولايات المتحدة": "United States", "إيران": "Iran",
    "نيجيريا": "Nigeria", "غانا": "Ghana", "إثيوبيا": "Ethiopia", "كينيا": "Kenya",
    "أفغانستان": "Afghanistan", "إندونيسيا": "Indonesia", "تركيا": "Turkey", "مالي": "Mali",
}


def _parse_date(s):
    if not s:
        return None
    if isinstance(s, (datetime, date)):
        return s if isinstance(s, date) else s.date()
    s = str(s).strip().split("T")[0].split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def fmt_date(s):
    d = _parse_date(s)
    return d.strftime("%d/%m/%Y") if d else (s or "")


def fmt_money(v):
    if v in (None, ""):
        return ""
    try:
        f = float(v)
        return str(int(f)) if f.is_integer() else f"{f:.3f}".rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        return str(v)


def resolve_contract_template(emp, company=None, signatory=None, project=None, contract_date=None):
    """يرجّع 22+ حقل للعقد. الأسماء القديمة (auth_name, sponsor, ...) مدعومة للتوافق مع القوالب القديمة."""
    emp = emp or {}
    company = company or {}
    signatory = signatory or {}
    project = project or {}
    start = contract_date or emp.get("dateOfHire") or ""
    d = _parse_date(start)
    nat = emp.get("nationality") or ""
    prof = emp.get("profession") or ""
    ctx = {
        "employee_name": emp.get("name") or "",
        "employee_name_en": emp.get("nameEn") or "",
        "nationality": nat,
        "nationality_en": emp.get("nationalityEn") or NATIONALITY_EN.get(nat, ""),
        "civil_id": emp.get("id") or "",
        "profession": prof,
        "profession_en": emp.get("professionEn") or PROFESSION_EN.get(prof, ""),
        "salary": fmt_money(emp.get("salary")),
        "housing_amount": fmt_money(emp.get("housingAmount")),
        "start_date": fmt_date(start),
        "day_name": AR_DAYS[d.weekday()] if d else "",
        "day_name_en": EN_DAYS[d.weekday()] if d else "",
        "passport_no": emp.get("passportNo") or "",
        "residency_exp": fmt_date(emp.get("residencyExp")),
        "file_number": project.get("fileNumber") or emp.get("fileNo") or company.get("mainFileNumber") or "",
        "company_name": company.get("nameAr") or "",
        "company_name_en": company.get("nameEn") or "",
        "labor_office": project.get("laborOffice") or company.get("laborOffice") or "",
        "project_name": project.get("nameAr") or "",
        "auth_name": signatory.get("nameAr") or "",
        "auth_name_en": signatory.get("nameEn") or "",
        "auth_civil_id": signatory.get("civilId") or "",
        "today": datetime.now().strftime("%d/%m/%Y"),
    }
    # أسماء قديمة من نظام العقود الآلي
    ctx["sponsor"] = ctx["company_name"]
    ctx["status"] = emp.get("transferNote") or ""
    ctx["end_date"] = fmt_date(emp.get("workPermitExp"))
    return ctx


# ---------------------------------------------------------------------------
# الاستبدال داخل الفقرات مع الحفاظ على التنسيق
# ---------------------------------------------------------------------------
def _replace_in_paragraph(p, repl_func, pattern):
    runs = p.runs
    if not runs:
        return
    full = "".join(r.text for r in runs)
    if not pattern.search(full):
        return
    # خريطة: موضع كل حرف ← (رقم الـ run، الموضع داخله)
    spans = []
    pos = 0
    for i, r in enumerate(runs):
        spans.append((pos, pos + len(r.text), i))
        pos += len(r.text)

    def run_at(ch):
        for s, e, i in spans:
            if s <= ch < e:
                return i, ch - s
        return len(runs) - 1, len(runs[-1].text)

    matches = list(pattern.finditer(full))
    texts = [r.text for r in runs]
    # من الآخر للأول عشان المواضع ما تتغيرش
    for m in reversed(matches):
        new = repl_func(m)
        if new is None:
            continue
        si, so = run_at(m.start())
        ei, eo = run_at(m.end() - 1)
        eo += 1
        if si == ei:
            texts[si] = texts[si][:so] + new + texts[si][eo:]
        else:
            texts[si] = texts[si][:so] + new
            for k in range(si + 1, ei):
                texts[k] = ""
            texts[ei] = texts[ei][eo:]
    for r, t in zip(runs, texts):
        if r.text != t:
            r.text = t


def _iter_paragraphs(doc):
    def from_container(c):
        for p in c.paragraphs:
            yield p
        for t in c.tables:
            for row in t.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    yield from from_container(cell)

    yield from from_container(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try:
                yield from from_container(part)
            except Exception:
                pass


def fill_docx_template(template_path, context):
    doc = Document(template_path)

    def braced(m):
        return str(context.get(m.group(1), ""))

    for p in _iter_paragraphs(doc):
        _replace_in_paragraph(p, braced, PLACEHOLDER_RE)
    _fill_merge_fields(doc, context)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fill_merge_fields(doc, context):
    """MERGEFIELD بسيطة: <w:fldSimple w:instr=' MERGEFIELD name '> → نص."""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for fld in doc.element.body.iter(ns + "fldSimple"):
        instr = fld.get(ns + "instr") or ""
        m = re.search(r"MERGEFIELD\s+([A-Za-z0-9_]+)", instr)
        if not m:
            continue
        val = str(context.get(m.group(1), ""))
        ts = list(fld.iter(ns + "t"))
        if ts:
            ts[0].text = val
            for t in ts[1:]:
                t.text = ""


def replace_literals(template_path, out_path, mapping):
    """أداة لتحويل عقد حقيقي إلى قالب: تستبدل نصوص حرفية بـ {{ field }} مع الحفاظ على التنسيق."""
    doc = Document(template_path)
    keys = sorted(mapping.keys(), key=len, reverse=True)
    pattern = re.compile("|".join(re.escape(k) for k in keys))
    for p in _iter_paragraphs(doc):
        _replace_in_paragraph(p, lambda m: mapping[m.group(0)], pattern)
    doc.save(out_path)


def list_placeholders(template_path):
    doc = Document(template_path)
    found = set()
    for p in _iter_paragraphs(doc):
        found.update(PLACEHOLDER_RE.findall("".join(r.text for r in p.runs)))
    return sorted(found)


# ---------------------------------------------------------------------------
# معاينة HTML و PDF
# ---------------------------------------------------------------------------
def _esc(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _para_html(p):
    parts = []
    for r in p.runs:
        t = _esc(r.text)
        if not t:
            continue
        if r.bold:
            t = f"<b>{t}</b>"
        if r.underline:
            t = f"<u>{t}</u>"
        parts.append(t)
    txt = "".join(parts).replace("\n", "<br>")
    align = {1: "center", 2: "left", 3: "justify"}.get(
        p.alignment if isinstance(p.alignment, int) else (p.alignment.real if p.alignment is not None else -1), ""
    )
    style = f' style="text-align:{align}"' if align else ""
    return f"<p{style}>{txt or '&nbsp;'}</p>"


def docx_to_html(data):
    doc = Document(io.BytesIO(data))
    body = doc.element.body
    out = []
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            out.append(_para_html(Paragraph(child, doc)))
        elif tag == "tbl":
            t = Table(child, doc)
            out.append('<table class="docx-table">')
            for row in t.rows:
                out.append("<tr>")
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    rtl = any("؀" <= ch <= "ۿ" for ch in cell.text[:200])
                    out.append(f'<td dir="{"rtl" if rtl else "ltr"}">')
                    for p in cell.paragraphs:
                        out.append(_para_html(p))
                    out.append("</td>")
                out.append("</tr>")
            out.append("</table>")
    return "\n".join(out)


def soffice_path():
    return shutil.which("libreoffice") or shutil.which("soffice")


def docx_to_pdf(data):
    exe = soffice_path()
    if not exe:
        # محاولة docx2pdf على ويندوز (يحتاج MS Word)
        try:
            from docx2pdf import convert  # type: ignore
        except ImportError:
            return None
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "c.docx")
            dst = os.path.join(d, "c.pdf")
            open(src, "wb").write(data)
            try:
                convert(src, dst)
                return open(dst, "rb").read()
            except Exception:
                return None
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "contract.docx")
        open(src, "wb").write(data)
        try:
            subprocess.run(
                [exe, "--headless", "--convert-to", "pdf", src, "--outdir", d],
                check=True, timeout=90, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            return open(os.path.join(d, "contract.pdf"), "rb").read()
        except Exception:
            return None
